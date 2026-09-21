from __future__ import annotations

from pathlib import Path

import openpyxl


class XlsxReader:
    """Doc tat ca sheet co du lieu trong XLSX."""

    @staticmethod
    def _tao_contract() -> dict:
        return {
            "status": "success",
            "file_type": "xlsx",
            "pdf_mode": None,
            "ocr_da_su_dung": False,
            "du_lieu": [],
            "canh_bao": [],
            "loi": None,
        }

    @staticmethod
    def _tao_header_duy_nhat(headers):
        ket_qua = []
        da_gap = {}

        for i, value in enumerate(
            headers,
            start=1,
        ):
            if (
                value is None
                or not str(value).strip()
            ):
                ten = f"Column_{i}"
            else:
                ten = str(value).strip()

            if ten not in da_gap:
                da_gap[ten] = 0
                ket_qua.append(ten)
                continue

            da_gap[ten] += 1

            ten_moi = (
                f"{ten}_{da_gap[ten]}"
            )

            while ten_moi in da_gap:
                da_gap[ten] += 1

                ten_moi = (
                    f"{ten}_{da_gap[ten]}"
                )

            da_gap[ten_moi] = 0
            ket_qua.append(ten_moi)

        return ket_qua

    def read(self, file_path: str | Path) -> dict:
        path = Path(file_path)
        result = self._tao_contract()

        workbook = None

        try:
            workbook = openpyxl.load_workbook(
                path,
                data_only=True,
                read_only=True,
            )

            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]

                rows = list(
                    sheet.iter_rows(
                        values_only=True
                    )
                )

                # Tim dong dau tien co du lieu
                header_index = -1

                for i, row in enumerate(rows):
                    if any(
                        cell is not None
                        and str(cell).strip()
                        for cell in row
                    ):
                        header_index = i
                        break

                if header_index == -1:
                    result["canh_bao"].append(
                        f"Sheet '{sheet_name}' rong, bo qua."
                    )
                    continue

                so_cot = max(
                    len(row)
                    for row in rows[header_index:]
                )

                raw_headers = list(
                    rows[header_index]
                )

                raw_headers += [
                    None
                ] * (
                    so_cot - len(raw_headers)
                )

                headers = (
                    self._tao_header_duy_nhat(
                        raw_headers
                    )
                )

                data_rows = []

                for row in rows[
                    header_index + 1:
                ]:
                    if not any(
                        cell is not None
                        and str(cell).strip()
                        for cell in row
                    ):
                        continue

                    row = list(row)

                    row += [
                        None
                    ] * (
                        so_cot - len(row)
                    )

                    row_dict = {}

                    for i, header in enumerate(
                        headers
                    ):
                        value = row[i]

                        if value is None:
                            value = ""
                        else:
                            value = (
                                str(value).strip()
                            )

                        row_dict[header] = value

                    data_rows.append(
                        row_dict
                    )

                result["du_lieu"].append(
                    {
                        "sheet_name": sheet_name,
                        "headers": headers,
                        "rows": data_rows,
                    }
                )

            if not result["du_lieu"]:
                result["status"] = "error"
                result["loi"] = "XLSX_NO_DATA"

        except FileNotFoundError:
            result["status"] = "error"
            result["loi"] = "FILE_NOT_FOUND"

        except Exception as e:
            result["status"] = "error"
            result["loi"] = (
                f"XLSX_READ_ERROR: {str(e)}"
            )

        finally:
            if workbook is not None:
                try:
                    workbook.close()
                except Exception:
                    pass

        return result